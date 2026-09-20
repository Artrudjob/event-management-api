import pytest
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from rest_framework import status
from events.constants import XLSX_HEADERS, XlsxFieldConstants
from events.models import Event, Venue

def _build_xlsx(rows: list[dict]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active

    for column, header in enumerate(XLSX_HEADERS, start=1):
        sheet.cell(row=1, column=column, value=header)

    for row_index, row in enumerate(rows, start=2):
        for column, field in enumerate(XLSX_HEADERS.values(), start=1):
            value = row[field]
            if isinstance(value, datetime) and timezone.is_aware(value):
                value = timezone.localtime(value).replace(tzinfo=None)
            sheet.cell(row=row_index, column=column, value=value)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    return stream

@pytest.mark.django_db
class TestEventImportExistingVenue:
    """Ожидается, что импорт в уже существующее место обновляет его, а не падает по unique."""

    @patch("events.services.event_import_service.fetch_weather_for_venue.delay")
    def test_update_exist_venue(self, _delay, superuser_client, venue):
        now = timezone.now()
        new_latitude = Decimal("59.934280")
        new_longitude = Decimal("30.335100")
        row = {
            XlsxFieldConstants.NAME: "Импортируемое мероприятие",
            XlsxFieldConstants.DESCRIPTION: "Описание",
            XlsxFieldConstants.PUBLICATION_AT: now + timedelta(days=1),
            XlsxFieldConstants.STARTS_AT: now + timedelta(days=2),
            XlsxFieldConstants.ENDS_AT: now + timedelta(days=3),
            XlsxFieldConstants.VENUE_NAME: venue.name,
            XlsxFieldConstants.LATITUDE: new_latitude,
            XlsxFieldConstants.LONGITUDE: new_longitude,
            XlsxFieldConstants.RATING: 10,
        }
        xlsx_file = SimpleUploadedFile(
            "events.xlsx",
            _build_xlsx([row]).read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = superuser_client.post(
            reverse("event-import-events"),
            {"file": xlsx_file},
            format="multipart",
        )
        venue.refresh_from_db()

        assert response.status_code == status.HTTP_201_CREATED
        assert response.data["created"] == 1
        assert Event.objects.filter(name=row[XlsxFieldConstants.NAME], venue=venue).exists()
        assert venue.latitude == new_latitude
        assert venue.longitude == new_longitude
        assert Venue.objects.filter(name=venue.name).count() == 1
