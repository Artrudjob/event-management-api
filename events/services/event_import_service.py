from zipfile import BadZipFile
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from events.constants import XLSX_HEADERS, XlsxFieldConstants
from events.models import Event, Venue
from events.serializers import EventImportSerializer
from weather.tasks import fetch_weather_for_venue


class EventXlsxParseError(Exception):
    def __init__(self, message: str):
        super().__init__(message)
        self.message = message

class EventImportService:
    def import_xlsx(self, file, author) -> dict:
        result = self._validate_xlsx(file)
        if result["errors"]:
            return result

        self._save_rows(result["rows"], author)

        return result

    def _parse_xlsx(self, file) -> list[dict]:
        workbook = None
        try:
            workbook = self._load_workbook(file)
            sheet = workbook.active
            if sheet is None:
                raise EventXlsxParseError("В файле нет листа с данными.")

            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                raise EventXlsxParseError("Файл пустой.")

            column_map = self._build_map_headers(header_row)
            parsed_rows = []

            for index, row in enumerate(rows, start=2):
                values = self._get_row_values(column_map, row)
                if self._is_empty(values):
                    continue

                values["row"] = index
                parsed_rows.append(values)

            return parsed_rows
        finally:
            if workbook is not None:
                workbook.close()

    def _validate_xlsx(self, file) -> dict:
        parsed_rows = self._parse_xlsx(file)
        rows = []
        errors = []

        for item in parsed_rows:
            row_number = item["row"]
            payload = {key: value for key, value in item.items() if key != "row"}
            serializer = EventImportSerializer(data=payload)

            if serializer.is_valid():
                rows.append({**serializer.validated_data, "row": row_number})
            else:
                errors.extend(self._flatten_serializer_errors(row_number, serializer.errors))

        return {"rows": rows, "errors": errors}

    def _save_rows(self, rows: list[dict], author) -> None:
        venue_ids = set()

        with transaction.atomic():
            for row in rows:
                venue = self._upsert_venue(row)
                venue_ids.add(venue.pk)
                Event.objects.create(
                    name=row[XlsxFieldConstants.NAME],
                    description=row[XlsxFieldConstants.DESCRIPTION] or "",
                    publication_at=row[XlsxFieldConstants.PUBLICATION_AT],
                    starts_at=row[XlsxFieldConstants.STARTS_AT],
                    ends_at=row[XlsxFieldConstants.ENDS_AT],
                    rating=row[XlsxFieldConstants.RATING],
                    status=row[XlsxFieldConstants.STATUS],
                    venue=venue,
                    author=author,
                )

            transaction.on_commit(
                lambda: self._queue_weather_fetch(venue_ids)
            )

    def _upsert_venue(self, row: dict) -> Venue:
        venue, created = Venue.objects.get_or_create(
            name=row[XlsxFieldConstants.VENUE_NAME],
            defaults={
                "latitude": row[XlsxFieldConstants.LATITUDE],
                "longitude": row[XlsxFieldConstants.LONGITUDE],
            },
        )
        if created:
            return venue

        venue.latitude = row[XlsxFieldConstants.LATITUDE]
        venue.longitude = row[XlsxFieldConstants.LONGITUDE]
        venue.save(update_fields=["latitude", "longitude", "update_time"])

        return venue

    def _queue_weather_fetch(self, venue_ids: set[int]) -> None:
        for venue_id in venue_ids:
            fetch_weather_for_venue.delay(venue_id)

    def _flatten_serializer_errors(self, row_number: int, serializer_errors) -> list[dict]:
        result = []

        for field, messages in serializer_errors.items():
            if not isinstance(messages, (list, tuple)):
                messages = [messages]
            for message in messages:
                result.append({
                    "row": row_number,
                    "field": field,
                    "message": str(message),
                })

        return result

    def _load_workbook(self, file):
        try:
            return load_workbook(file, read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, OSError) as exc:
            raise EventXlsxParseError("Не удалось прочитать xlsx-файл.") from exc

    def _build_map_headers(self, header_row) -> dict[int, str]:
        column_map = {}
        seen_fields = set()

        for index, header in enumerate(header_row):
            if header is None or str(header).strip() == "":
                continue

            header_name = str(header).strip().lower()
            field = XLSX_HEADERS.get(header_name)
            if field is None:
                continue

            if field in seen_fields:
                raise EventXlsxParseError(f"Колонка «{header_name}» указана больше одного раза.")

            column_map[index] = field
            seen_fields.add(field)

        missing = [
            header
            for header, field in XLSX_HEADERS.items()
            if field not in seen_fields
        ]
        if missing:
            missing_list = ", ".join(missing)
            raise EventXlsxParseError(f"В файле нет обязательных колонок: {missing_list}.")

        return column_map

    def _get_row_values(self, column_map, row) -> dict:
        values = {}
        for index, field in column_map.items():
            cell = row[index] if index < len(row) else None
            values[field] = cell

        return values

    def _is_empty(self, values: dict) -> bool:
        return all(value is None or str(value).strip() == "" for value in values.values())
