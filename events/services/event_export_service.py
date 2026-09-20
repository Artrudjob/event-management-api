from io import BytesIO
from django.utils import timezone
from openpyxl import Workbook
from events.constants import XLSX_HEADERS, XlsxFieldConstants
from events.models import Event

class EventExportService:
    FIELD_MAPPERS = {
        XlsxFieldConstants.NAME: lambda event: event.name,
        XlsxFieldConstants.DESCRIPTION: lambda event: event.description,
        XlsxFieldConstants.PUBLICATION_AT: lambda event: EventExportService._prepare_datetime(event.publication_at),
        XlsxFieldConstants.STARTS_AT: lambda event: EventExportService._prepare_datetime(event.starts_at),
        XlsxFieldConstants.ENDS_AT: lambda event: EventExportService._prepare_datetime(event.ends_at),
        XlsxFieldConstants.VENUE_NAME: lambda event: event.venue.name,
        XlsxFieldConstants.LATITUDE: lambda event: event.venue.latitude,
        XlsxFieldConstants.LONGITUDE: lambda event: event.venue.longitude,
        XlsxFieldConstants.RATING: lambda event: event.rating,
    }

    def export_xlsx(self, events) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        self._set_header(sheet)

        for row_index, event in enumerate(events, start=2):
            self._set_row(sheet, row_index, event)

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)

        return stream

    def _set_header(self, sheet) -> None:
        for column, header in enumerate(XLSX_HEADERS, start=1):
            sheet.cell(row=1, column=column, value=header)

    def _set_row(self, sheet, row_index: int, event: Event) -> None:
        for column, field in enumerate(XLSX_HEADERS.values(), start=1):
            sheet.cell(row=row_index, column=column, value=self.FIELD_MAPPERS[field](event))

    @staticmethod
    def _prepare_datetime(datetime):
        return timezone.localtime(datetime).replace(tzinfo=None)
